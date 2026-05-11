"""将论文初稿.md转换为符合华工本科毕设格式的Word文档"""
import re
import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_line_spacing(paragraph, spacing_pt=None, line_spacing=1.5):
    """设置段落行距"""
    pf = paragraph.paragraph_format
    if line_spacing:
        pf.line_spacing = line_spacing
    if spacing_pt is not None:
        pf.space_before = Pt(spacing_pt)
        pf.space_after = Pt(spacing_pt)

def set_run_font(run, font_name_cn, font_name_en, size_pt, bold=False):
    """设置run的字体"""
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = font_name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), font_name_cn)
    rFonts.set(qn('w:ascii'), font_name_en)
    rFonts.set(qn('w:hAnsi'), font_name_en)
    rPr.insert(0, rFonts)

def add_paragraph_with_style(doc, text, font_cn='宋体', font_en='Times New Roman',
                              size=12, bold=False, align=None, spacing_before=0,
                              spacing_after=0, first_line_indent=None, line_spacing=1.5):
    """添加带格式的段落"""
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    set_line_spacing(para, line_spacing=line_spacing)
    pf = para.paragraph_format
    pf.space_before = Pt(spacing_before)
    pf.space_after = Pt(spacing_after)
    if first_line_indent:
        pf.first_line_indent = first_line_indent

    run = para.add_run(text)
    set_run_font(run, font_cn, font_en, size, bold)
    return para

def add_heading_formatted(doc, text, level, number_text=None):
    """按照模板格式添加标题"""
    if level == 'chapter':  # 章标题：黑体 小二号(18pt) 居中 单倍行距 段前段后0.5行
        display = text
        if number_text:
            display = f"{number_text} {text}"
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_line_spacing(para, line_spacing=1.0)
        pf = para.paragraph_format
        pf.space_before = Pt(6)   # 0.5行 ≈ 6pt
        pf.space_after = Pt(6)
        run = para.add_run(display)
        set_run_font(run, '黑体', 'Times New Roman', 18, bold=True)
        return para

    elif level == 'section':  # 一级标题：黑体 小三号(15pt) 居左 单倍行距 段前段后0.5行
        display = text
        if number_text:
            display = f"{number_text} {text}"
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_line_spacing(para, line_spacing=1.0)
        pf = para.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        run = para.add_run(display)
        set_run_font(run, '黑体', 'Times New Roman', 15, bold=True)
        return para

    elif level == 'subsection':  # 二级标题：黑体 四号(14pt) 居左 单倍行距 段前段后0.5行
        display = text
        if number_text:
            display = f"{number_text} {text}"
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_line_spacing(para, line_spacing=1.0)
        pf = para.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        run = para.add_run(display)
        set_run_font(run, '黑体', 'Times New Roman', 14, bold=True)
        return para

    elif level == 'subsubsection':  # 三级标题：黑体 小四号(12pt) 居左
        display = text
        if number_text:
            display = f"{number_text} {text}"
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_line_spacing(para, line_spacing=1.0)
        pf = para.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        run = para.add_run(display)
        set_run_font(run, '黑体', 'Times New Roman', 12, bold=True)
        return para

def add_body_text(doc, text):
    """添加正文段落：宋体/Times New Roman 小四号(12pt) 1.5倍行距 首行缩进2字符"""
    para = doc.add_paragraph()
    set_line_spacing(para, line_spacing=1.5)
    pf = para.paragraph_format
    pf.first_line_indent = Pt(24)  # 约2个汉字宽度

    # 处理中文和英文/数字混合，分别设置字体
    # 简化处理：整段设置，中文优先
    run = para.add_run(text)
    set_run_font(run, '宋体', 'Times New Roman', 12)
    return para

def add_md_table(doc, table_lines):
    """将markdown表格行列表转换为Word表格"""
    if len(table_lines) < 2:
        return
    # 解析表格
    def parse_row(line):
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        return cells
    header = parse_row(table_lines[0])
    # 跳过分隔行 (|---|---|)
    data_lines = [parse_row(l) for l in table_lines[1:] if '---' in l or not l.startswith('|')]
    data_lines = [parse_row(l) for l in table_lines if not re.match(r'^[\|\s\-:]+$', l)]
    if not data_lines:
        return
    header = data_lines[0]
    rows = data_lines[1:]
    ncols = len(header)
    table = doc.add_table(rows=1 + len(rows), cols=ncols, style='Table Grid')
    table.autofit = True
    # 表头
    for j, cell_text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(cell_text)
        set_run_font(run, '黑体', 'Times New Roman', 10.5, bold=True)
    # 数据行
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j < ncols:
                cell = table.rows[i + 1].cells[j]
                cell.text = ''
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(cell_text)
                set_run_font(run, '宋体', 'Times New Roman', 10.5)
    # 表后空行
    blank = doc.add_paragraph()
    set_line_spacing(blank, line_spacing=1.0)
    blank.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text):
    """添加代码块"""
    for line in code_text.strip().split('\n'):
        para = doc.add_paragraph()
        set_line_spacing(para, line_spacing=1.0)
        pf = para.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        run = para.add_run(line)
        set_run_font(run, '等线', 'Consolas', 10)
    # 代码块后空一行
    blank = doc.add_paragraph()
    set_line_spacing(blank, line_spacing=1.0)
    blank.paragraph_format.space_after = Pt(6)

def add_table_caption(doc, text):
    """表标题：宋体五号(10.5pt) 居中"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(para, line_spacing=1.0)
    pf = para.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    run = para.add_run(text)
    set_run_font(run, '宋体', 'Times New Roman', 10.5, bold=True)
    return para

def add_figure_caption(doc, text):
    """图标题：宋体五号(10.5pt) 居中"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(para, line_spacing=1.0)
    pf = para.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    run = para.add_run(text)
    set_run_font(run, '宋体', 'Times New Roman', 10.5)
    return para

def add_image_figure(doc, img_path, caption_text=""):
    """插入图片并添加图标题，居中显示"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_line_spacing(para, line_spacing=1.0)
    pf = para.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    pf.first_line_indent = Pt(0)
    run = para.add_run()
    if os.path.exists(img_path):
        run.add_picture(img_path, width=Inches(5.0))
    else:
        run_text = para.add_run(f"[图片缺失: {img_path}]")
        set_run_font(run_text, '宋体', 'Times New Roman', 10.5)
    if caption_text:
        add_figure_caption(doc, caption_text)

def add_page_break(doc):
    """添加分页符"""
    doc.add_page_break()

def parse_chapter_number(heading_text):
    """从标题文本中提取章节号和标题名"""
    # 匹配: ## 第1章 绪论 或 ## 第一章 绪论
    patterns = [
        r'^第(\d+)章\s+(.+)',
        r'^第([一二三四五六七八九十]+)章\s+(.+)',
    ]
    for p in patterns:
        m = re.match(p, heading_text)
        if m:
            return m.group(1), m.group(2)
    return None, heading_text

def parse_section_number(heading_text):
    """解析节号"""
    # 匹配: 1.1 标题 或 1 标题
    m = re.match(r'^(\d+(?:\.\d+)*)\s+(.+)', heading_text)
    if m:
        return m.group(1), m.group(2)
    return None, heading_text

# Chinese number mapping
CN_NUMS = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7}

def cn_to_int(s):
    return CN_NUMS.get(s, int(s) if s.isdigit() else 1)


def generate_thesis(md_path, output_path):
    """主函数：读取markdown并生成格式化的Word文档"""
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 页面设置
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # ==================== 封面 ====================
    add_page_break(doc)
    for _ in range(4):
        doc.add_paragraph()

    add_paragraph_with_style(doc, '本科毕业设计（论文）', font_cn='黑体', font_en='Times New Roman',
                             size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                             spacing_before=12, spacing_after=24, line_spacing=1.0)

    add_paragraph_with_style(doc, '基于大语言模型的医疗自然语言查询系统设计与实现',
                             font_cn='黑体', font_en='Times New Roman',
                             size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                             line_spacing=1.0)

    for _ in range(6):
        doc.add_paragraph()

    # 封面信息表（简化为段落）
    info_items = [
        '学院：____________________',
        '专业：____________________',
        '学生姓名：________________',
        '学生学号：________________',
        '指导教师：________________',
        '提交日期：________________',
    ]
    for item in info_items:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_line_spacing(para, line_spacing=2.0)
        run = para.add_run(item)
        set_run_font(run, '宋体', 'Times New Roman', 14)

    add_page_break(doc)

    # ==================== 原创性声明 ====================
    add_paragraph_with_style(doc, '学位论文原创性声明',
                             font_cn='黑体', font_en='Times New Roman',
                             size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                             spacing_before=6, spacing_after=6, line_spacing=1.0)

    declaration = (
        '本人郑重声明：所呈交的论文是本人在导师的指导下独立进行研究所取得的研究成果。'
        '除了文中特别加以标注引用的内容外，本论文不包含任何其他个人或集体已经发表或撰写的成果作品。'
        '对本文的研究做出重要贡献的个人和集体，均已在文中以明确方式标明。'
        '本人完全意识到本声明的法律后果由本人承担。'
    )
    add_paragraph_with_style(doc, declaration, font_cn='宋体', font_en='Times New Roman',
                             size=14, line_spacing=1.5, first_line_indent=Cm(0.74))

    para = doc.add_paragraph()
    set_line_spacing(para, line_spacing=1.5)
    run = para.add_run('作者签名：                       日期：    年   月   日')
    set_run_font(run, '宋体', 'Times New Roman', 14)

    add_page_break(doc)

    # ==================== 版权使用授权书 ====================
    add_paragraph_with_style(doc, '学位论文版权使用授权书',
                             font_cn='黑体', font_en='Times New Roman',
                             size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                             spacing_before=6, spacing_after=6, line_spacing=1.0)

    auth_text = (
        '本学位论文作者完全了解学校有关保留、使用学位论文的规定，即：学校有权保存'
        '并向国家有关部门或机构送交论文的复印件和电子版，允许学位论文被查阅；学校可以'
        '公布学位论文的全部或部分内容，可以允许采用影印、缩印或其它复制手段保存、汇编'
        '学位论文。本人电子文档的内容和纸质论文的内容相一致。'
    )
    add_paragraph_with_style(doc, auth_text, font_cn='宋体', font_en='Times New Roman',
                             size=14, line_spacing=1.5, first_line_indent=Cm(0.74))

    para = doc.add_paragraph()
    set_line_spacing(para, line_spacing=1.5)
    run = para.add_run('作者签名：                        日期：    年   月   日')
    set_run_font(run, '宋体', 'Times New Roman', 14)

    para = doc.add_paragraph()
    set_line_spacing(para, line_spacing=1.5)
    run = para.add_run('指导教师签名：                    日期：    年   月   日')
    set_run_font(run, '宋体', 'Times New Roman', 14)

    add_page_break(doc)

    # ==================== 中文摘要 ====================
    add_paragraph_with_style(doc, '摘  要',
                             font_cn='黑体', font_en='Times New Roman',
                             size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                             spacing_before=6, spacing_after=6, line_spacing=1.0)

    abstract_cn = (
        '随着医疗信息化的快速发展，医院信息系统积累了海量的临床数据。这些数据存储在关系型数据库中，'
        '查询需要编写SQL语句，对不具备编程技能的医护人员构成了较高的使用门槛。本文设计并实现了一个'
        '基于大语言模型的医疗自然语言查询系统，用户只需输入中文自然语言问题，系统即可自动将其转换'
        '为SQL语句并执行查询。系统采用前后端分离架构，后端基于FastAPI框架与PyMySQL实现数据库交互'
        '与业务逻辑编排，前端基于Vue 3与Element Plus构建交互界面，大模型服务接入阿里云通义千问'
        '（Qwen-Plus）API完成自然语言到SQL的转换。系统的核心工作包括：动态获取数据库表结构以构建'
        '上下文化Prompt；采用基于白名单模式的双层SQL安全校验机制，确保仅允许SELECT类查询操作；'
        '实现查询历史持久化、数据库结构浏览等辅助功能。测试结果表明，系统能够有效支持单表查询、'
        '聚合统计、条件筛选等多类医疗查询场景，SQL生成准确率良好，端到端响应时间可满足日常使用需求。'
    )
    add_paragraph_with_style(doc, abstract_cn, font_cn='宋体', font_en='Times New Roman',
                             size=12, line_spacing=1.5, first_line_indent=Cm(0.74))
    doc.add_paragraph()  # 空一行

    # 关键词
    para = doc.add_paragraph()
    set_line_spacing(para, line_spacing=1.5)
    pf = para.paragraph_format
    pf.first_line_indent = Cm(0.74)
    kw_label = para.add_run('关键词：')
    set_run_font(kw_label, '黑体', 'Times New Roman', 12, bold=True)
    kw_text = para.add_run('自然语言转SQL；大语言模型；医疗信息系统；Prompt Engineering；FastAPI')
    set_run_font(kw_text, '宋体', 'Times New Roman', 12)

    add_page_break(doc)

    # ==================== 英文摘要 ====================
    add_paragraph_with_style(doc, 'Abstract',
                             font_cn='Times New Roman', font_en='Times New Roman',
                             size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                             spacing_before=6, spacing_after=6, line_spacing=1.0)

    abstract_en = (
        'With the rapid development of medical informatization, hospital information systems '
        'have accumulated massive amounts of clinical data. These data are typically stored in '
        'relational databases, and querying requires writing SQL statements, which poses a high '
        'barrier for medical staff without programming skills. This paper designs and implements '
        'a medical natural language query system based on large language models (LLMs). Users '
        'input natural language questions in Chinese, and the system automatically converts them '
        'into SQL statements and executes the queries. The system adopts a front-end and back-end '
        'separation architecture: the back-end uses FastAPI and PyMySQL for database interaction '
        'and business logic orchestration; the front-end uses Vue 3 and Element Plus for the user '
        'interface; the LLM service connects to Alibaba Cloud Qwen-Plus API via the OpenAI-compatible '
        'SDK to perform natural language to SQL conversion. Core contributions include dynamic '
        'database schema extraction for contextualized prompt construction, a whitelist-based '
        'two-layer SQL safety validation mechanism ensuring only SELECT queries are permitted, '
        'and auxiliary features including query history persistence and database schema browsing. '
        'Test results demonstrate that the system effectively supports common medical query scenarios '
        'including single-table queries, aggregate statistics, and conditional filtering, with '
        'favorable SQL generation accuracy and end-to-end response times meeting daily usage requirements.'
    )
    add_paragraph_with_style(doc, abstract_en, font_cn='Times New Roman', font_en='Times New Roman',
                             size=12, line_spacing=1.5, first_line_indent=Cm(0.74))
    doc.add_paragraph()

    para = doc.add_paragraph()
    set_line_spacing(para, line_spacing=1.5)
    pf = para.paragraph_format
    pf.first_line_indent = Cm(0.74)
    kw_label = para.add_run('Keywords: ')
    set_run_font(kw_label, 'Times New Roman', 'Times New Roman', 12, bold=True)
    kw_text = para.add_run('NL2SQL; Large Language Model; Medical Information System; Prompt Engineering; FastAPI')
    set_run_font(kw_text, 'Times New Roman', 'Times New Roman', 12)

    add_page_break(doc)

    # ==================== 目录占位 ====================
    add_paragraph_with_style(doc, '目  录',
                             font_cn='黑体', font_en='Times New Roman',
                             size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                             spacing_before=6, spacing_after=12, line_spacing=1.0)
    add_paragraph_with_style(doc, '（目录请在Word中使用"引用→目录"自动生成）',
                             font_cn='宋体', font_en='Times New Roman',
                             size=12, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5)
    add_page_break(doc)

    # ==================== 正文内容 ====================
    # 读取markdown文件
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # 找到正文开始位置（第1章 绪论）
    body_match = re.search(r'##\s+第1章\s+绪论', content)
    if not body_match:
        body_match = re.search(r'##\s+第一章\s+绪论', content)

    if body_match:
        body_content = content[body_match.start():]

        # 将正文分为行处理
        lines = body_content.split('\n')
        i = 0
        current_chapter_num = 0
        section_counter = {}
        subsection_counter = {}

        while i < len(lines):
            line = lines[i].strip()

            # 跳过空行
            if not line:
                i += 1
                continue

            # 匹配章标题: ## 第X章 标题
            chapter_match = re.match(r'^##\s+第(\d+)章\s+(.+)', line)
            if chapter_match:
                ch_num = int(chapter_match.group(1))
                ch_title = chapter_match.group(2)
                current_chapter_num = ch_num
                section_counter = {}
                subsection_counter = {}

                # 第7章改为"结论"
                if ch_num == 7:
                    display_title = '结论'
                else:
                    cn_nums_rev = {1: '一', 2: '二', 3: '三', 4: '四', 5: '五', 6: '六'}
                    display_title = f"第{cn_nums_rev.get(ch_num, str(ch_num))}章 {ch_title}"

                add_heading_formatted(doc, display_title, 'chapter')
                i += 1
                continue

            # 匹配一级标题: ### X.X 标题
            section_match = re.match(r'^###\s+(\d+)\.(\d+)\s+(.+)', line)
            if section_match:
                sec_major = int(section_match.group(1))
                sec_minor = int(section_match.group(2))
                sec_title = section_match.group(3)

                sec_key = sec_major
                if sec_key not in section_counter:
                    section_counter[sec_key] = 0
                section_counter[sec_key] += 1
                subsection_counter = {}

                # 使用简化编号（模板格式：章内序号）
                display_num = str(section_counter[sec_key])
                add_heading_formatted(doc, sec_title, 'section', number_text=display_num)
                i += 1
                continue

            # 匹配二级标题: #### X.X.X 标题
            subsection_match = re.match(r'^####\s+(\d+)\.(\d+)\.(\d+)\s+(.+)', line)
            if subsection_match:
                sub_title = subsection_match.group(4)
                sec_major = int(subsection_match.group(1))
                sub_key = int(subsection_match.group(3))
                if sub_key not in subsection_counter:
                    subsection_counter[sub_key] = 0
                subsection_counter[sub_key] += 1

                add_heading_formatted(doc, sub_title, 'subsection', number_text=str(subsection_counter[sub_key]))
                i += 1
                continue

            # 匹配代码块
            code_match = re.match(r'^```(\w*)$', line)
            if code_match:
                code_lines = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith('```'):
                    code_lines.append(lines[i])
                    i += 1
                add_code_block(doc, '\n'.join(code_lines))
                i += 1
                continue

            # 匹配表格行 — 收集连续表格行作为一个表格
            if line.startswith('|'):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    table_lines.append(lines[i].strip())
                    i += 1
                add_md_table(doc, table_lines)
                continue

            # 匹配一级/二级/三级列表
            list_match = re.match(r'^[-*]\s+(.+)', line)
            numbered_match = re.match(r'^\((\d+)\)\s+(.+)', line)
            bold_section = re.match(r'^\*\*(.+?)\*\*[：:]\s*(.*)', line)

            if list_match:
                text = list_match.group(1)
                # 处理列表中的粗体
                text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
                add_body_text(doc, text)
                i += 1
                continue

            if numbered_match:
                text = f"（{numbered_match.group(1)}）{numbered_match.group(2)}"
                add_body_text(doc, text)
                i += 1
                continue

            if bold_section:
                text = f"{bold_section.group(1)}：{bold_section.group(2)}"
                add_body_text(doc, text)
                i += 1
                continue

            # 跳过markdown水平线、图片链接等
            if line in ('---', '---',):
                i += 1
                continue

            # 跳过引用标记
            if line.startswith('>'):
                i += 1
                continue

            # 图片链接: ![caption](path) 或后跟正文
            img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)(.*)', line)
            if img_match:
                caption = img_match.group(1)
                img_path = img_match.group(2)
                remaining = img_match.group(3).strip()
                if not os.path.isabs(img_path):
                    img_path = os.path.join(os.path.dirname(md_path), img_path)
                add_image_figure(doc, img_path, caption)
                if remaining:
                    add_body_text(doc, remaining)
                i += 1
                continue

            # 清理行内的markdown标记
            clean_line = line
            clean_line = re.sub(r'\*\*(.+?)\*\*', r'\1', clean_line)  # 粗体
            clean_line = re.sub(r'\*(.+?)\*', r'\1', clean_line)      # 斜体
            clean_line = re.sub(r'`([^`]+)`', r'\1', clean_line)      # 内联代码
            clean_line = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', clean_line)  # 链接

            # 跳过纯标题符号
            if re.match(r'^#{1,6}\s', clean_line):
                i += 1
                continue

            # 正文段落
            if clean_line and len(clean_line) > 5:
                add_body_text(doc, clean_line)

            i += 1

    # ==================== 结论（从第7章内容提取） ====================
    # 已在上面通过章节匹配处理

    # ==================== 参考文献 ====================
    add_page_break(doc)
    add_heading_formatted(doc, '参考文献', 'chapter')

    refs = [
        '[1] Vaswani A, Shazeer N, Parmar N, et al. Attention is all you need[C]. Advances in Neural Information Processing Systems, 2017: 5998-6008.',
        '[2] Brown T B, Mann B, Ryder N, et al. Language models are few-shot learners[C]. Advances in Neural Information Processing Systems, 2020: 1877-1901.',
        '[3] Ouyang L, Wu J, Jiang X, et al. Training language models to follow instructions with human feedback[C]. Advances in Neural Information Processing Systems, 2022: 27730-27744.',
        '[4] Zhong V, Xiong C, Socher R. Seq2SQL: Generating structured queries from natural language using reinforcement learning[J]. arXiv preprint arXiv:1709.00103, 2017.',
        '[5] Yu T, Zhang R, Yang K, et al. Spider: A large-scale human-labeled dataset for complex and cross-domain semantic parsing and text-to-SQL task[C]. Proceedings of EMNLP, 2018: 3911-3921.',
        '[6] Scholak T, Schucher N, Bahdanau D. PICARD: Parsing incrementally for constrained auto-regressive decoding from language models[C]. Proceedings of EMNLP, 2021: 9895-9901.',
        '[7] Pourreza M, Rafiei D. DIN-SQL: Decomposed in-context learning of text-to-SQL with self-correction[J]. arXiv preprint arXiv:2304.11015, 2023.',
        '[8] Bai J, Bai S, Chu Y, et al. Qwen technical report[J]. arXiv preprint arXiv:2309.16609, 2023.',
        '[9] 杨宝华, 王磊. 基于深度学习的NL2SQL技术研究综述[J]. 计算机应用研究, 2022, 39(6): 1601-1608.',
        '[10] 李俊, 张志强. 医疗信息系统中的数据查询技术综述[J]. 医学信息学杂志, 2021, 42(3): 28-33.',
    ]

    for ref in refs:
        para = doc.add_paragraph()
        set_line_spacing(para, line_spacing=1.5)
        run = para.add_run(ref)
        set_run_font(run, '宋体', 'Times New Roman', 12)

    # ==================== 致谢 ====================
    add_page_break(doc)
    add_heading_formatted(doc, '致谢', 'chapter')

    thanks_text = (
        '本论文（设计）在导师的悉心指导下完成。从选题方向的确定、技术方案的讨论到论文的撰写与修改，'
        '导师给予了大量宝贵的指导和建议。在此，谨向导师表示崇高的敬意和衷心的感谢。'
    )
    add_body_text(doc, thanks_text)

    thanks_text2 = (
        '感谢在毕业设计过程中给予帮助的各位老师和同学。在系统开发和测试阶段，他们提出了许多有价值的'
        '意见和建议，帮助我发现了潜在的问题并加以改进。'
    )
    add_body_text(doc, thanks_text2)

    thanks_text3 = (
        '感谢开源社区提供的优秀工具和框架，包括FastAPI、Vue.js、Element Plus、PyMySQL等，'
        '它们极大降低了系统开发的复杂度。感谢阿里云DashScope平台提供的通义千问大模型API服务，'
        '为本系统的核心功能提供了强大的技术支撑。'
    )
    add_body_text(doc, thanks_text3)

    thanks_text4 = (
        '感谢我的家人和朋友在我求学期间的长期理解与支持，他们的鼓励是我坚持完成学业的重要动力。'
    )
    add_body_text(doc, thanks_text4)

    # ==================== 保存 ====================
    doc.save(output_path)
    print(f'论文已生成: {output_path}')


if __name__ == '__main__':
    import os
    base = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(base, '论文初稿.md')
    output_path = os.path.join(base, '论文初稿_格式版.docx')
    generate_thesis(md_path, output_path)
